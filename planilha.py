import openpyxl

import mysql.connector
conexao = mysql.connector.connect(
            host='localhost',
            database='certificados', 
            user='root', 
            password='')




# Carregar o arquivo Excel
book = openpyxl.load_workbook('AlunosTurmas.xlsx')
ws = book.active
# Selecionar a planilha desejada
sheet = book['1 - TDS3IN-A']  # Substitua pelo nome da sua planilha
ws.tables

for table in ws.tables.values():
    print(table)

# Iterar sobre as linhas e colunas para pegar os valores
num_min = str(8)
num_max = str(8)
min = 'A' + num_min
max = 'I' + num_max
cell_range = sheet[min+':'+ max]
contador = 0
alunos = []
while contador < 20:
    for row in cell_range:
        for cell in row:
            alunos.append(cell.value)
    cursor = conexao.cursor()
    cursor.execute("SELECT * FROM resp")
print(alunos[1]) 