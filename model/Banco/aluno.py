import string
import model.Banco.conexao as conexao
import openpyxl
from docx import Document

#Função responsável por deletar todos os alunos cadastrados
def deletar_todos_alunos(alunos):
    try:
        conexao_bd = conexao.iniciar_conexao()
        for aluno in alunos:
            query = 'DELETE FROM aluno WHERE id = %s;'
            parametro = [aluno[0]]
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
        return True
    except:
        return False
    finally:
        cursor_bd.close()
        conexao_bd.close()


#Função responsável pelo cadastro das informações dos alunos por meio de planilhas
def cadastrar_turma_alunos(planilha):
    try:
        wb = openpyxl.load_workbook(planilha)

        alunos = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                if all(cell is None or str(cell).strip() == '' for cell in row):
                    continue
                alunos.append(row)


        wb.close()

        for aluno in alunos:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO aluno (nome_completo, nome_mae, nome_pai, cpf, municipio, estado, data_nascimento, rg, orgao_expedidor, turma) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s);"
            parametro = (aluno[0], aluno[1], aluno[2], aluno[3], aluno[4], aluno[5], aluno[6], aluno[7], aluno[8], aluno[9])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        return False
    

#Função responsável pelo cadastro das informações do curso por meio de planilhas
def cadastrar_curso(planilha):
    try:
        wb = openpyxl.load_workbook(planilha)

        cursos = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                if all(cell is None or str(cell).strip() == '' for cell in row):
                    continue
                cursos.append(row)

        for curso in cursos:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO curso_tecnico (nome_curso, eixo_tecnologico, perfil_profissional, carga_horaria_total, data_conclusao) VALUES (%s, %s, %s, %s, %s);"
            parametro = (curso[0], curso[1], curso[2], curso[3], curso[4])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        print("Erro ao inserir curso:", curso)
        return False

#Função responsável pelo cadastro das disciplina da base comum dos alunos por meio de planilhas
def cadastrar_turma_disciplinas_base_comum(planilha):
    try:
        wb = openpyxl.load_workbook(planilha)

        materias_bc = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                if all(cell is None or str(cell).strip() == '' for cell in row):
                    continue
                materias_bc.append(row)

        wb.close()

        for bc in materias_bc:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO base_comum_disciplinas (nome_disciplina, ano_escolar, ano_realizacao, carga_horaria) VALUES (%s, %s, %s, %s);"
            parametro = (bc[0], bc[1], bc[2], bc[3])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
            print('PEGOUU', materias_bc)
        return True
    except:
        return False


#Função responsável pelo cadastro das disciplina das trilhas dos alunos por meio de planilhas
def cadastrar_turma_disciplinas_trilha(planilha, sheet):
    try:
        wb = openpyxl.load_workbook(planilha)
        sheet = wb[sheet]
        materias_trilha = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                materias_trilha.append(row)

        wb.close()

        for trilha in materias_trilha:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO trilha_disciplinas (nome_disciplina, ano_escolar, ano_realizacao, carga_horaria, frequencia, tuma) VALUES (%s, %s, %s, %s, %s, %s);"
            parametro = (trilha[2], trilha[3], trilha[4], trilha[5], trilha[6], trilha[7])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        return False


#Função responsável pelo cadastro das disciplina das eletivas dos alunos por meio de planilhas
def cadastrar_turma_disciplinas_eletiva(planilha, sheet):
    try:
        wb = openpyxl.load_workbook(planilha)
        sheet = wb[sheet]
        materias_eletiva = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                materias_eletiva.append(row)

        wb.close()

        for eletiva in materias_eletiva:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO eletiva_disciplinas (nome_disciplina, ano_escolar, ano_realizacao, carga_horaria, frequencia, tuma) VALUES (%s, %s, %s, %s, %s, %s);"
            parametro = (eletiva[2], eletiva[3], eletiva[4], eletiva[5], eletiva[6], eletiva[7])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        return False
    

#Função responsável pelo cadastro das disciplinas do curso por meio de planilhas
def cadastrar_turma_disciplinas_curso(planilha, sheet):
    try:
        wb = openpyxl.load_workbook(planilha)
        sheet = wb[sheet]
        materias_curso = []

        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row = 2, values_only=True):
                materias_curso.append(row)

        wb.close()

        for curso in materias_curso:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO curso_tecnico_disciplinas (nome_disciplina, ano_escolar, ano_realizacao, carga_horaria, frequencia) VALUES (%s, %s, %s, %s, %s, %s);"
            parametro = (curso[2], curso[3], curso[4], curso[5], curso[6])
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        return False


#INCOMPLETA
#Função responsável por cadastrar um aluno no próprio site
def cadastrar_aluno(nome_aluno, nome_mae, nome_pai, cpf_aluno, municipio_aluno, estado_aluno, nascimento_aluno, rg_aluno, orgao_aluno, turma_aluno):
    try:
            conexao_bd = conexao.iniciar_conexao()
            query = "INSERT INTO aluno (nome_completo, nome_mae, nome_pai, cpf, municipio, estado, data_nascimento, rg, orgao_expedidor, turma) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s);"
            parametro = (nome_aluno, nome_mae, nome_pai, cpf_aluno, municipio_aluno, estado_aluno, nascimento_aluno, rg_aluno, orgao_aluno, turma_aluno)
            cursor_bd = conexao_bd.cursor()
            cursor_bd.execute(query, parametro)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
            return True
    except:
        return False


#Função responsável por listar os alunos cadastrados em suas respectivas turmas
def listar_alunos(seletor):
    lista = []
    conexao_bd = conexao.iniciar_conexao()
    if seletor == None or seletor == "*":
        query = "SELECT * FROM aluno"

        cursor_bd = conexao_bd.cursor()
        cursor_bd.execute(query)
        for aluno in cursor_bd:
            lista.append(aluno)
        cursor_bd.close()
        conexao_bd.close()
        return lista
    
    elif seletor == "3TDSA":
        query = "SELECT * FROM aluno WHERE turma = '3TDSA';"

        cursor_bd = conexao_bd.cursor()
        cursor_bd.execute(query)
        for aluno in cursor_bd:
            lista.append(aluno)
        cursor_bd.close()
        conexao_bd.close()
        return lista
    
    elif seletor == "3TDSB":
        query = "SELECT * FROM aluno WHERE turma = '3TDSB';"

        cursor_bd = conexao_bd.cursor()
        cursor_bd.execute(query)
        for aluno in cursor_bd:
            lista.append(aluno)
        cursor_bd.close()
        conexao_bd.close()

        if lista != None:
            return lista
    
    elif seletor == "3MKTA":
        query = "SELECT * FROM aluno WHERE turma = '3MKTA';"

        cursor_bd = conexao_bd.cursor()
        cursor_bd.execute(query)
        for aluno in cursor_bd:
            lista.append(aluno)
        cursor_bd.close()
        conexao_bd.close()
        return lista
    
    elif seletor == "3MKTB":
        query = "SELECT * FROM aluno WHERE turma = '3MKTB';"

        cursor_bd = conexao_bd.cursor()
        cursor_bd.execute(query)
        for aluno in cursor_bd:
            lista.append(aluno)
        cursor_bd.close()
        conexao_bd.close()
        return lista


#Função responsável por exibir informações do aluno
def exibir_aluno(id_aluno, id_curso):
    conexao_bd = conexao.iniciar_conexao()
    query = "SELECT * FROM aluno WHERE id = " + str(id_aluno) +";"
    cursor_bd = conexao_bd.cursor()

    cursor_bd.execute(query)
    info_aluno = cursor_bd.fetchone()
    cursor_bd.close()
    conexao_bd.close()
    return info_aluno 


#Função responsável por editar as informações dos alunos que já estão cadastrados
def editar_aluno(nome_aluno, nome_mae, nome_pai, cpf_aluno, municipio_aluno, estado_aluno, rg_aluno, turma_aluno, nome_curso, eixo_tecnologico):
    conexao_bd = conexao.iniciar_conexao()

    query_aluno = "UPDATE aluno SET nome_completo = %s, nome_mae = %s, nome_pai = %s, cpf = %s, municipio = %s, estado = %s, rg = %s turma = %s WHERE id = %s;"
    parametros_aluno = (nome_aluno, nome_mae, nome_pai, cpf_aluno, municipio_aluno, estado_aluno, rg_aluno, turma_aluno)

    query_curso = "UPDATE curso_tecnico SET nome_curso = %s, eixo tecnologico = %s;"
    parametros_curso = (nome_curso, eixo_tecnologico)

    cursor_bd = conexao_bd.cursor()

    cursor_bd.execute(query_aluno, query_curso, parametros_aluno, parametros_curso)
    conexao_bd.commit()
    cursor_bd.close()
    conexao_bd.close()


#Função responsável por deletar um aluno
def deletar_aluno(id_aluno):
    conexao_bd = conexao.iniciar_conexao()
    query = 'DELETE FROM bdalunos WHERE id = %s;'
    parametro = [id_aluno]
    cursor_bd = conexao_bd.cursor()

    cursor_bd.execute(query, parametro)
    conexao_bd.commit()
    cursor_bd.close()
    conexao_bd.close()