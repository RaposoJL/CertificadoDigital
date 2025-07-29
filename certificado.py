import string
import model.Banco.conexao as conexao
import locale
from datetime import datetime
from docx import Document
import os 
import shutil

#Função que identifica a data de emissão do certificado, guardando ela no banco de dados
def data_hoje(data_emissao):
    try:
        hoje = datetime.now()
        locale.setlocale(locale.LC_TIME, 'pt_BR.UTF-8')
        data_marcada = hoje.strftime('%d de %B de %Y')

        for data in data_marcada:
            conexao_bd = conexao.iniciar_conexao()
            cursor_bd = conexao_bd.cursor()
            marcar = "INSERT INTO certificado (data_emissao) VALUES (%s);"
            parametros = (data_emissao) 
            cursor_bd.execute(marcar, parametros)
            conexao_bd.commit()
            cursor_bd.close()
            conexao_bd.close()
        return True
    except:
        return False
    

#INCOMPLETA
#Função responsável por gerar somente um certificado
def gerar_certificado(documento, id_aluno, id_certificado, id_curso, id_bc_disciplinas, id_curso_disciplinas, id_trilha_disciplinas, id_eletiva_disciplinas):
    conexao_bd = conexao.iniciar_conexao()
    cursor_bd = conexao_bd.cursor()

    cursor_bd.execute("SELECT * FROM aluno WHERE id = %s", str(id_aluno,))
    info_aluno = cursor_bd.fetchone()

    cursor_bd.execute("SELECT * FROM certificado WHERE id = %s", str(id_certificado,))
    info_certificado = cursor_bd.fetchone()

    cursor_bd.execute("SELECT * FROM curso_tecnico WHERE id = %s", (id_curso,))
    info_curso = cursor_bd.fetchone()

    cursor_bd.execute("SELECT nome_disciplina, carga_horaria FROM base_comum_disciplinas WHERE id = %s", (id_bc_disciplinas,))
    info_bc_disciplinas = cursor_bd.fetchall()

    cursor_bd.execute("SELECT nome_disciplina, carga_horaria FROM curso_tecnico_disciplinas WHERE id = %s", (id_curso_disciplinas,))
    info_curso_disciplinas = cursor_bd.fetchall()

    cursor_bd.execute("SELECT nome_disciplina, carga_horaria FROM trilha_disciplinas WHERE id = %s", (id_trilha_disciplinas,))
    info_trilha_disciplinas = cursor_bd.fetchall()

    cursor_bd.execute("SELECT nome_disciplina, carga_horaria FROM eletiva_disciplinas WHERE id = %s", (id_eletiva_disciplinas,))
    info_eletiva_disciplinas = cursor_bd.fetchall()

    cursor_bd.close()
    conexao_bd.close()

    doc = Document(documento)

    substituicoes = {
        '#nome_completo#':  info_aluno[1],
        '#nome_mae#':  info_aluno[2],
        '#nome_pai#':  info_aluno[3],
        '#municipio#':  info_aluno[4],
        '#estado#':  info_aluno[5],
        '#data_nascimento#':  info_aluno[6],
        '#rg#':  info_aluno[7],
        '#orgao_expedidor#':  info_aluno[8],
        '#cpf#':  info_aluno[9],
        '#data_conclusao#': info_curso[6],
        '#nome_curso#': info_curso[2],
        '#eixo_tecnologico#': info_curso[3],
        '#data_emissao#': info_certificado[2],
        '#perfil_profissional#': info_curso[4],
        '#cht#': info_curso[5]
    }

    todas_disciplinas = info_bc_disciplinas + info_curso_disciplinas + info_trilha_disciplinas + info_eletiva_disciplinas

    for i, (nome_disciplina, carga_horaria) in enumerate(todas_disciplinas, start=1):
        substituicoes[f'#disciplinas{i}#'] = nome_disciplina
        substituicoes[f'#ch{i}#'] = f"{carga_horaria}"

    for par in doc.paragraphs:
        for run in par.runs:
            for palavra_antiga, palavra_nova in substituicoes.items():
                if palavra_antiga in run.text:
                    run.text = run.text.replace(palavra_antiga, palavra_nova)

    tabela = doc.tables[0]

    for tabela in doc.tables:
        for row in tabela.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for run in par.runs:
                        for palavra_antiga, palavra_nova in substituicoes.items():
                            if palavra_antiga in run.text:
                                run.text = run.text.replace(palavra_antiga, palavra_nova)

    pasta_certificados = os.path.abspath("static/Certificados/" + str(info_aluno[0]) + "-" + info_aluno[1] +".docx")
    doc.save(pasta_certificados)


#INCOMPLETA
#Função responsável por gerar todos os certificados
def gerar_todos_certificados(alunos, documento):
    for aluno in alunos:
        gerar_certificado(documento, aluno[0])
    
    shutil.make_archive("static/Certificado", 'zip', "static/Certificados/")
    return "Certificado"


#Função responsável por realizar a busca por aluno
def pesquisar_aluno(nome_aluno):
    conexao_bd = conexao.iniciar_conexao()
    cursor_bd = conexao_bd.cursor()

    cadastrar = f'SELECT * FROM aluno WHERE nome_completo LIKE "{nome_aluno}%"'

    cursor_bd.execute(cadastrar)
    info_aluno = cursor_bd.fetchall()
    cursor_bd.close()
    conexao_bd.close()
    return info_aluno


#Função responsável por exibir certificados que já foram criados
def certificados(certificados_criados):
    lista = []
    conexao_bd = conexao.iniciar_conexao()
    cursor_bd = conexao_bd.cursor()
    
    for id_certificado in certificados_criados:
        query = "SELECT * FROM certificado WHERE id = " + str(id_certificado)
        cursor_bd.execute(query)
        lista.append(cursor_bd.fetchone())

    cursor_bd.close()
    conexao_bd.close()
    return lista